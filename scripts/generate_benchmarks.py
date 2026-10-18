"""Write ONLY synthetic benchmark files; no private source data is used."""
import argparse
import csv
from datetime import datetime
import io
from pathlib import Path
import re
import zipfile

from openpyxl import Workbook


def generate(output):
    out = Path(output)
    out.mkdir(parents=True, exist_ok=True)

    def write(name, headers, rows):
        with (out/name).open('w', encoding='utf-8', newline='') as handle:
            writer = csv.writer(handle, delimiter=';', lineterminator='\n')
            writer.writerow(headers)
            writer.writerows(rows)

    headers = ['ID', 'Tipo', 'Valor', 'Moeda', 'Pago', 'Data Lançamento', 'Competência', 'Data Vencimento', 'Data Pagamento', 'Categoria', 'Cliente_ID']
    clean = [
        ['001', 'Receita', '1200,50', 'BRL', 'Sim', '2026-01-01', '2026-01-01', '2026-01-15', '2026-01-14', 'Serviços', 'C01'],
        ['002', 'Despesa', '-300,20', 'BRL', 'Sim', '2026-01-02', '2026-01-01', '2026-01-10', '2026-01-10', 'Infraestrutura', 'C02'],
        ['003', 'Receita', '500,00', 'BRL', 'Não', '2026-01-03', '2026-01-01', '2026-02-15', '', 'Serviços', 'C03'],
        ['004', 'Despesa', '-100,00', 'BRL', 'Não', '2026-01-04', '2026-01-01', '2026-02-15', '', 'Licenças', 'C02'],
        ['005', 'Receita', '700,00', 'BRL', 'Sim', '2026-02-01', '2026-02-01', '2026-02-10', '2026-02-09', 'Serviços', 'C01'],
        ['006', 'Despesa', '-250,30', 'BRL', 'Sim', '2026-02-02', '2026-02-01', '2026-02-10', '2026-02-10', 'Infraestrutura', 'C02'],
        ['007', 'Receita', '900,00', 'BRL', 'Sim', '2026-02-03', '2026-02-01', '2026-02-20', '2026-02-20', 'Serviços', 'C03'],
        ['008', 'Despesa', '-80,00', 'BRL', 'Sim', '2026-02-04', '2026-02-01', '2026-02-20', '2026-02-20', 'Licenças', 'C02'],
    ]
    write('finance_clean.csv', headers, clean)
    dirty = [list(r) for r in clean]
    dirty[0][8] = ''
    dirty[1][2] = 'erro'
    dirty[2][3] = 'USD'
    dirty[3][4] = '?'
    dirty[4][5] = '2026-02-30'
    dirty[6][9] = 'servicos'
    dirty.append(list(dirty[5]))
    write('finance_dirty.csv', headers, dirty)
    write('finance_ambiguous.csv', ['Valor', 'Data', 'Pago'], [
        ['1,234', '01/02/2026', '?'], ['2.345', '02/03/2026', 'talvez'], ['100', '03/04/2026', '']])
    write('regional_br.csv', ['Valor', 'Data', 'Moeda'], [
        ['1.234', '01/02/2026', 'BRL'], ['(1.234,56)', '13/02/2026', 'BRL'], ['0,10', '2026-02-14', 'BRL']])
    write('regional_us.csv', ['Valor', 'Data', 'Moeda'], [
        ['(1,234.56)', '02/13/2026', 'USD'], ['0.10', '02/14/2026', 'USD']])
    write('legitimate_duplicates.csv', ['Tipo', 'Valor', 'Moeda', 'Data Lançamento'], [
        ['Despesa', '-50,00', 'BRL', '2026-01-15'], ['Despesa', '-50,00', 'BRL', '2026-01-15'],
        ['Receita', '200,00', 'BRL', '2026-01-15']])
    write('transfers_reversals.csv', ['ID', 'Tipo', 'Valor', 'Moeda', 'Pago', 'Data Pagamento', 'Conta'], [
        ['T1', 'Transferência', '-500,00', 'BRL', 'Sim', '2026-01-15', 'Banco A'],
        ['T2', 'Transferência', '500,00', 'BRL', 'Sim', '2026-01-15', 'Banco B'],
        ['T3', 'Estorno', '-100,00', 'BRL', 'Sim', '2026-01-16', 'Banco B'],
        ['T4', 'Reembolso', '100,00', 'BRL', 'Sim', '2026-01-17', 'Banco A'],
        ['T5', 'Receita', '1000,00', 'BRL', 'Cancelado', '', 'Banco A']])
    write('projects_missing_dates.csv', ['Projeto', 'Atividade', 'Data Atualização', 'Data Conclusão'], [
        ['P01', 'Planejar', '2026-01-15', '2026-01-20'], ['P01', 'Executar', '2026-01-16', '']])
    wb = Workbook()
    wb.active.title = 'Movimentos'
    wb.active.append(headers)
    for row in clean:
        wb.active.append(row)
    clients = wb.create_sheet('Clientes')
    clients.append(['Cliente_ID', 'Nome'])
    for row in [('C01', 'Cliente sintético A'), ('C02', 'Cliente sintético B'), ('C02', 'Cadastro duplicado')]:
        clients.append(row)
    categories = wb.create_sheet('Categorias')
    categories.append(['Categoria_ID', 'Categoria'])
    categories.append(['01', 'Serviços'])
    wb.properties.creator = 'PowerBI Autopilot synthetic benchmark'
    wb.properties.created = datetime(2026, 1, 1)
    wb.properties.modified = datetime(2026, 1, 1)
    buffer = io.BytesIO()
    wb.save(buffer)
    # XLSX is ZIP: remove timestamps, including openpyxl's modified property.
    with zipfile.ZipFile(buffer) as src, zipfile.ZipFile(out/'finance_related.xlsx', 'w', compression=zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            content = src.read(name)
            if name == 'docProps/core.xml':
                content = re.sub(rb'(<dcterms:modified[^>]*>).*?(</dcterms:modified>)', rb'\g<1>2026-01-01T00:00:00Z\2', content)
            info = zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, content)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', default='tests/fixtures')
    generate(parser.parse_args().output)
