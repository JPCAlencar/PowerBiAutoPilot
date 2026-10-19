"""Independent QA probes for local intake (V0.2).

Every source is synthesized in a temporary directory; nothing from the repository
fixtures is reused.
"""
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from powerbi_autopilot.intake import InputError, ReadOptions, read_source
from powerbi_autopilot.profiling import profile_table


class IntakeCase(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp(prefix='qa-intake-'))
        self.addCleanup(shutil.rmtree, self.dir, True)

    def csv(self, name, text, encoding='utf-8'):
        path = self.dir / name
        path.write_text(text, encoding=encoding, newline='')
        return path

    def binary(self, name, payload):
        path = self.dir / name
        path.write_bytes(payload)
        return path

    def workbook(self, name, rows, sheet='Dados', hidden=None, empty=None):
        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        for row in rows:
            ws.append(row)
        if hidden:
            extra = wb.create_sheet('Oculta')
            for row in hidden:
                extra.append(row)
            extra.sheet_state = 'hidden'
        if empty:
            wb.create_sheet('Vazia')
        path = self.dir / name
        wb.save(path)
        return path


class CsvIntakeTests(IntakeCase):
    def test_bom_is_stripped_by_default(self):
        path = self.binary('bom.csv', 'ID;Valor\n1;10,00\n'.encode('utf-8-sig'))
        t = read_source(path)[0]
        self.assertEqual(list(t.frame.columns), ['ID', 'Valor'])
        self.assertEqual(t.source['encoding'], 'utf-8-sig')
        self.assertEqual(t.source['delimiter'], ';')
        self.assertEqual(t.row_numbers, [2])

    def test_duplicate_and_blank_headers_are_disambiguated(self):
        path = self.csv('headers.csv', 'ID;ID;;Valor\n1;2;3;10\n')
        t = read_source(path)[0]
        self.assertEqual(list(t.frame.columns), ['ID', 'ID__2', 'column_3', 'Valor'])
        found = {(i['code'], i['severity']) for i in t.issues}
        self.assertIn(('duplicate_header', 'WARNING'), found)
        self.assertIn(('blank_header', 'WARNING'), found)
        # CSV keeps the raw empty string; Excel blanks arrive as None.
        self.assertEqual([h['original'] for h in t.source['headers']], ['ID', 'ID', '', 'Valor'])
        self.assertEqual([h['name'] for h in t.source['headers']], ['ID', 'ID__2', 'column_3', 'Valor'])

    def test_quoted_multiline_record_keeps_identity(self):
        path = self.csv('multi.csv', 'ID;Obs\n1;"linha um\nlinha dois"\n2;ok\n')
        t = read_source(path, ReadOptions(sep=';'))[0]
        self.assertEqual(len(t.frame), 2)
        self.assertEqual(t.frame['Obs'].tolist(), ['linha um\nlinha dois', 'ok'])
        self.assertEqual(len(set(t.row_numbers)), 2)

    def test_ragged_row_is_rejected(self):
        path = self.csv('ragged.csv', 'ID;Valor\n1;10\n2;20;30\n')
        # Delimiter sniffing fails first, so the message does not name the row.
        with self.assertRaises(InputError) as ctx:
            read_source(path)
        self.assertIn('Delimitador incerto ou linhas irregulares', str(ctx.exception))
        # With an explicit separator the offending line is identified.
        with self.assertRaises(InputError) as explicit:
            read_source(path, ReadOptions(sep=';'))
        self.assertIn('linha 3', str(explicit.exception))

    def test_footer_total_row_is_preserved_and_flagged_by_profiling(self):
        # QA-F14: intake preserves Raw; profiling raises a bounded hypothesis,
        # never silently promoting the row to a confirmed detail transaction.
        path = self.csv('footer.csv', 'ID;Valor\n1;10,00\nTotal;10,00\n')
        t = read_source(path)[0]
        self.assertEqual(len(t.frame), 2)
        self.assertEqual(t.row_numbers, [2, 3])
        self.assertEqual(t.issues, [])
        self.assertEqual(t.frame['ID'].tolist(), ['1', 'Total'])
        p = profile_table(t, 'pt-BR')
        event = next(i for i in p['issues'] if i['code'] == 'footer_row_suspected')
        self.assertEqual(event['severity'], 'WARNING')
        self.assertEqual(event['source_rows'], [3])
        self.assertEqual(event['evidence_kind'], 'inference')
        self.assertEqual(p['grain']['status'], 'unknown')

    def test_single_column_csv_with_decimal_comma_is_preserved(self):
        # QA-F16: decimal commas in the body are not structural separators when
        # the logical header has one field. The assumption remains explicit.
        path = self.csv('uma_coluna.csv', 'Valor\n10,00\n20,00\n')
        t = read_source(path)[0]
        self.assertEqual(t.frame['Valor'].tolist(), ['10,00', '20,00'])
        self.assertEqual(t.row_numbers, [2, 3])
        self.assertIsNone(t.source['delimiter'])
        self.assertIn('single_column_csv', {i['code'] for i in t.issues})
        # The same file is read correctly when the separator is explicit.
        t = read_source(path, ReadOptions(sep=';'))[0]
        self.assertEqual(t.frame['Valor'].tolist(), ['10,00', '20,00'])
        self.assertEqual(t.source['delimiter'], ';')

    def test_unsupported_payloads_are_refused(self):
        cases = {
            'nul.csv': b'ID;Valor\n1;\x00\n',
            'empty.csv': b'   \n',
            'noheader.csv': b'\n',
            'binary.csv': b'\xff\xfe\x00A',
        }
        for name, payload in cases.items():
            with self.subTest(name=name):
                with self.assertRaises(InputError):
                    read_source(self.binary(name, payload))
        with self.assertRaises(InputError):
            read_source(self.csv('x.txt', 'ID\n1\n'))
        with self.assertRaises(InputError):
            read_source(self.csv('sep.csv', 'ID;Valor\n1;2\n', ), ReadOptions(sheet='Aba'))
        with self.assertRaises(InputError):
            read_source(self.csv('opt.csv', 'ID;Valor\n1;2\n'), ReadOptions(sep=';;'))
        with self.assertRaises(InputError):
            read_source(self.csv('row.csv', 'ID;Valor\n1;2\n'), ReadOptions(header_row=0))

    def test_non_utf8_falls_back_to_cp1252_with_a_warning(self):
        path = self.binary('latin.csv', 'Situa\xe7\xe3o;Valor\nPago;10,00\n'.encode('cp1252'))
        t = read_source(path)[0]
        self.assertEqual(t.source['encoding'], 'cp1252')
        self.assertIn('Situação', list(t.frame.columns))
        self.assertEqual([i['code'] for i in t.issues], ['encoding_fallback'])
        # An explicit wrong encoding must not guess silently.
        with self.assertRaises(InputError):
            read_source(path, ReadOptions(encoding='utf-8'))


class ExcelIntakeTests(IntakeCase):
    def test_all_sheets_are_read_and_ids_are_stable(self):
        path = self.workbook('book.xlsx', [['ID', 'Valor'], ['1', '10,00']],
                             hidden=[['ID', 'X'], ['1', 'y']], empty=True)
        tables = read_source(path)
        self.assertEqual([t.source['sheet'] for t in tables], ['Dados', 'Oculta', 'Vazia'])
        ids = [t.id for t in tables]
        self.assertEqual(len(set(ids)), 3)
        self.assertTrue(all(i.startswith(tables[0].source['sha256'][:16]) for i in ids))
        self.assertEqual(len(tables[2].frame), 0)
        self.assertEqual(len(tables[2].frame.columns), 0)
        selected = read_source(path, ReadOptions(sheet='Oculta'))
        self.assertEqual([t.source['sheet'] for t in selected], ['Oculta'])
        with self.assertRaises(InputError):
            read_source(path, ReadOptions(sheet='Inexistente'))

    def test_sparse_sheet_assigns_positional_header(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Dados'
        ws['A1'] = 'ID'
        ws['B1'] = 'Valor'
        ws['C5'] = 'solto'
        path = self.dir / 'sparse.xlsx'
        wb.save(path)
        t = read_source(path)[0]
        self.assertEqual(list(t.frame.columns), ['ID', 'Valor', 'column_3'])
        self.assertEqual(len(t.frame), 4)
        self.assertEqual([i['code'] for i in t.issues], ['blank_header'])

    def test_header_row_beyond_data_reads_nothing_but_does_not_crash(self):
        path = self.workbook('short.xlsx', [['ID', 'Valor'], ['1', '10,00']])
        t = read_source(path, ReadOptions(header_row=10))[0]
        self.assertEqual(len(t.frame), 0)
        self.assertEqual(list(t.frame.columns), [])
        self.assertEqual(t.source['header_row'], 10)

    def test_formulas_are_never_evaluated(self):
        path = self.workbook('formula.xlsx', [['ID', 'Valor'], ['1', '=1+1']])
        t = read_source(path)[0]
        self.assertEqual(t.frame['Valor'].tolist(), [None])
        self.assertEqual([i['code'] for i in t.issues], ['formula_not_evaluated'])
        self.assertEqual(t.issues[0]['severity'], 'BLOCKER')
        self.assertEqual(t.issues[0]['source_rows'], [2])

    def test_qa_f05_chartsheet_must_not_hide_the_data_sheets(self):
        """QA-F05 (CONFIRMED BUG, MEDIUM): ``Workbook.sheetnames`` also lists chart
        sheets, and ``Chartsheet`` has no ``iter_rows``; the resulting
        AttributeError is converted into ``InputError('Planilha inválida,
        protegida ou ilegível')``. A perfectly valid workbook is rejected and the
        message blames corruption/protection. Expected: tabular sheets are read and
        the chart sheet is reported as non-tabular (or skipped with an INFO).
        """
        wb = Workbook()
        ws = wb.active
        ws.title = 'Dados'
        ws.append(['ID', 'Valor'])
        ws.append(['1', '10,00'])
        wb.create_chartsheet('Grafico')
        path = self.dir / 'chart.xlsx'
        wb.save(path)
        tables = read_source(path)
        self.assertEqual([t.source['sheet'] for t in tables], ['Dados'])
        self.assertEqual(len(tables[0].frame), 1)


if __name__ == '__main__':
    unittest.main()
