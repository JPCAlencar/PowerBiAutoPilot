"""V0.2.1 regressions: synthetic inputs, independent literal expectations."""
from copy import deepcopy
from decimal import Decimal
import hashlib
from pathlib import Path
import tempfile
import unittest

import pandas as pd

from powerbi_autopilot.contracts import Table
from powerbi_autopilot import financial, parsing, validation
from powerbi_autopilot.profiling import profile_table
from powerbi_autopilot.semantics import map_columns
from powerbi_autopilot.intake import InputError, ReadOptions, read_source
from powerbi_autopilot.engine import analyze
from powerbi_autopilot.planning import plan_transformations
from powerbi_autopilot.relationships import candidates


def table(data):
    frame = pd.DataFrame(data, dtype=object)
    return Table('synthetic', frame, {'file': 'synthetic.csv', 'sheet': None,
                                     'sha256': 'synthetic'}, list(range(2, len(frame) + 2)))


def diagnose(t, locale='pt-BR'):
    p = profile_table(t, locale)
    return financial.diagnose(t, p, map_columns(t, p, locale), locale)


class MagnitudeTests(unittest.TestCase):
    def test_zero_prefixed_grouping_is_refused_in_every_grouping_path(self):
        for raw, locale in [('0,500', 'en-US'), ('0.001', 'pt-BR'),
                            ('-US$ 0,500', 'en-US'), ('(R$ 0.500)', 'pt-BR'),
                            ('01,234', 'en-US'), ('00.123', 'pt-BR'),
                            ('0,123,456', None), ('0.123.456', None),
                            ('0,123.45', None), ('00.123,45', None)]:
            with self.subTest(raw=raw, locale=locale):
                result = parsing.number(raw, locale)
                self.assertEqual(result.status, 'invalid')
                self.assertIsNone(result.value)

    def test_fractional_decimals_and_legitimate_grouping_keep_exact_magnitude(self):
        for raw, locale, expected in [('0,500', 'pt-BR', '0.500'),
                                      ('0.001', 'en-US', '0.001'),
                                      ('1,234,567.89', 'en-US', '1234567.89'),
                                      ('1.234.567,89', 'pt-BR', '1234567.89'),
                                      ('1,234,567', None, '1234567'),
                                      ('0', 'pt-BR', '0')]:
            with self.subTest(raw=raw, locale=locale):
                self.assertEqual(parsing.number(raw, locale).value, Decimal(expected))
        self.assertEqual(parsing.number('0,500').status, 'ambiguous')
        self.assertEqual(parsing.number('0.500').status, 'ambiguous')

    def test_magnitude_error_cannot_enter_profile_or_cash_requirements(self):
        for raw, locale in [('0,500', 'en-US'), ('0.500', 'pt-BR')]:
            with self.subTest(locale=locale):
                t = table({'Valor': [raw], 'Moeda': ['BRL'], 'Tipo': ['Receita'],
                           'Status': ['Pago'], 'Data Pagamento': ['2026-01-05']})
                original = t.frame.copy(deep=True)
                p = profile_table(t, locale)
                self.assertIsNone(p['columns'][0]['number_distribution'])
                event = next(i for i in p['issues'] if i['code'] == 'invalid_number')
                self.assertEqual(event['source_rows'], [2])
                metric = next(m for m in diagnose(t, locale)['indicators'] if m['id'] == 'cash_received')
                self.assertIn('amount', metric['missing_requirements'])
                pd.testing.assert_frame_equal(t.frame, original)

    def test_reconciliation_refuses_equal_but_misinterpreted_amounts(self):
        for raw, locale in [('0,500', 'en-US'), ('0.500', 'pt-BR')]:
            with self.subTest(locale=locale):
                t = table({'Valor': [raw], 'Moeda': ['BRL']})
                r = validation.reconcile(t, deepcopy(t), ['Valor'], 'Moeda', locale)
                self.assertEqual(r['status'], 'failed')
                self.assertEqual(r['totals'], {'Valor': {}})
                self.assertIn('unreconciled_value', {i['code'] for i in r['issues']})


class SemanticSafetyTests(unittest.TestCase):
    def test_numeric_task_targets_do_not_activate_financial_status_audit(self):
        r = diagnose(table({'Tarefa': ['Entrega', 'Projeto'], 'Meta': ['10', '20'],
                            'Tipo': ['Marco', 'Atividade'],
                            'Status': ['Concluído', 'Em andamento']}))
        self.assertEqual(r['domain']['name'], 'unknown')
        self.assertFalse({'unknown_payment_status', 'unknown_movement_type'} &
                         {i['code'] for i in r['issues']})

    def test_negated_and_qualified_payment_flags_do_not_infer_polarity(self):
        for name in ('Não Pago', 'Nao Foi Pago', 'Não Liquidado', 'Valor Pago', 'Foi Pago'):
            with self.subTest(name=name):
                r = diagnose(table({'Valor': ['10', '20'], name: ['1', '0']}))
                self.assertEqual(r['roles'].get('payment_status', []), [])
                self.assertNotIn('paid_without_payment_date', {i['code'] for i in r['issues']})

    def test_generic_binary_status_never_declares_payment(self):
        for flags in (['1', '0'], [1, 0], [True, False], ['Sim', 'Não'], ['S', 'N']):
            with self.subTest(flags=flags):
                t = table({'Valor': ['100,00', '50,00'], 'Moeda': ['BRL'] * 2,
                           'Tipo': ['Receita', 'Despesa'], 'Status': flags})
                r = diagnose(t)
                self.assertEqual(r['roles'].get('payment_status', []), [])
                self.assertFalse({'unknown_payment_status', 'paid_without_payment_date'} &
                                 {i['code'] for i in r['issues']})

    def test_payment_named_flag_keeps_boolean_compatibility(self):
        for name in ('Pago', 'Paid', 'Quitado', 'Liquidado'):
            with self.subTest(name=name):
                r = diagnose(table({name: [True, False], 'Valor': ['100', '50'],
                                    'Moeda': ['BRL'] * 2, 'Tipo': ['Receita', 'Despesa'],
                                    'Data Pagamento': ['2026-01-01', None]}))
                self.assertEqual(r['roles'].get('payment_status'), [name])
                self.assertNotIn('paid_without_payment_date', {i['code'] for i in r['issues']})

    def test_mixed_generic_status_does_not_turn_a_flag_into_paid(self):
        r = diagnose(table({'Valor': ['100', '50'], 'Moeda': ['BRL'] * 2,
                            'Status': ['Pago', '1'],
                            'Data Pagamento': ['2026-01-01', None]}))
        self.assertNotIn('paid_without_payment_date', {i['code'] for i in r['issues']})
        event = next(i for i in r['issues'] if i['code'] == 'unknown_payment_status')
        self.assertEqual(event['source_rows'], [3])
        self.assertEqual(event['severity'], 'BLOCKER')

    def test_compound_headers_require_values_and_do_not_match_unrelated_names(self):
        for amount, payment_date in [('Valor Total', 'Data de Pagamento'),
                                     ('Vlr. Pgto', 'Data de Recebimento'),
                                     ('Valor Total', 'Dt Pagto')]:
            with self.subTest(amount=amount, payment_date=payment_date):
                t = table({amount: ['100', '50'], payment_date: ['2026-01-01'] * 2})
                roles = {s['role']: c['column'] for c in map_columns(t, profile_table(t))
                         for s in c['suggestions']}
                self.assertEqual(roles.get('amount'), amount)
                self.assertEqual(roles.get('payment_date'), payment_date)
                invalid = table({amount: ['texto'], payment_date: ['31/02/2026']})
                self.assertFalse(any(c['suggestions'] for c in map_columns(invalid, profile_table(invalid))))
        t = table({'Total de Itens': ['10'], 'Data de Cadastro': ['2026-01-01'],
                   'Desvalorização': ['10'], 'Status Ativo': ['1']})
        roles = {s['role'] for c in map_columns(t, profile_table(t)) for s in c['suggestions']}
        self.assertFalse({'amount', 'payment_date', 'payment_status'} & roles)

    def test_financial_unknown_fields_still_block_without_becoming_roles(self):
        r = diagnose(table({'Valor': ['100', '50'], 'Moeda': ['BRL'] * 2,
                            'Tipo': ['indefinido'] * 2, 'Status': ['Parcial', 'Baixado']}))
        self.assertEqual(r['roles'].get('movement_type', []), [])
        self.assertEqual(r['roles'].get('payment_status', []), [])
        by_code = {i['code']: i for i in r['issues']}
        self.assertEqual(by_code['unknown_movement_type']['source_rows'], [2, 3])
        self.assertEqual(by_code['unknown_payment_status']['source_rows'], [2, 3])
        self.assertNotIn('paid_without_payment_date', by_code)

    def test_invalid_compound_payment_date_still_blocks(self):
        r = diagnose(table({'Valor Total': ['100'], 'Tipo': ['Receita'],
                            'Status': ['Pago'], 'Moeda': ['BRL'],
                            'Data de Pagamento': ['31/02/2026']}))
        event = next(i for i in r['issues'] if i['code'] == 'paid_without_payment_date')
        self.assertEqual(event['column'], 'Data de Pagamento')
        self.assertEqual(event['source_rows'], [2])

    def test_abbreviated_amount_and_payment_dates_keep_conversion_diagnostics(self):
        t = table({'Vlr. Pgto': ['erro'], 'Dt Pagto': ['31/02/2026']})
        p = profile_table(t, 'pt-BR')
        findings = {(i['code'], i['column']) for i in p['issues']}
        self.assertIn(('invalid_number', 'Vlr. Pgto'), findings)
        self.assertIn(('invalid_date', 'Dt Pagto'), findings)

    def test_multiple_payment_dates_are_not_chosen_by_fallback(self):
        # Supply an additional valid rule rather than relying on an invented alias.
        from powerbi_autopilot.semantics import DEFAULT_RULES, SemanticRule
        t = table({'Valor': ['100'], 'Status': ['Pago'], 'Moeda': ['BRL'],
                   'Data Pagamento': ['2026-01-01'], 'Alternativa': ['2026-01-02']})
        p = profile_table(t)
        mapping = map_columns(t, p, rules=DEFAULT_RULES +
                              (SemanticRule('payment_date', ('alternativa',), 'date'),))
        r = financial.diagnose(t, p, mapping)
        event = next(i for i in r['issues'] if i['code'] == 'paid_without_payment_date')
        self.assertIsNone(event['column'])

    def test_em_aberto_is_pending_without_inventing_overdue(self):
        r = diagnose(table({'Valor': ['100'], 'Tipo': ['Receita'], 'Moeda': ['BRL'],
                            'Status': ['Em aberto'], 'Data Pagamento': ['2026-01-01']}))
        self.assertEqual(r['roles'].get('payment_status'), ['Status'])
        self.assertNotIn('unknown_payment_status', {i['code'] for i in r['issues']})
        overdue = next(m for m in r['indicators'] if m['id'] == 'overdue')
        self.assertIn('overdue_policy', overdue['missing_requirements'])
        self.assertNotIn('value', overdue)


class IntakeSafetyTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def csv(self, content):
        path = self.root / 'synthetic.csv'
        path.write_text(content, encoding='utf-8', newline='')
        return path

    def test_single_column_decimal_preserves_quotes_blank_rows_and_bytes(self):
        path = self.csv('Valor\n10,00\n"20,00"\n\n0,500\n')
        raw = path.read_bytes()
        t = read_source(path)[0]
        self.assertEqual(t.frame['Valor'].tolist(), ['10,00', '20,00', None, '0,500'])
        self.assertEqual(t.row_numbers, [2, 3, 4, 5])
        self.assertIsNone(t.source['delimiter'])
        self.assertEqual(t.source['sha256'], hashlib.sha256(raw).hexdigest())
        self.assertEqual(path.read_bytes(), raw)
        event = next(i for i in t.issues if i['code'] == 'single_column_csv')
        self.assertEqual(event['evidence_kind'], 'inference')

    def test_logical_header_and_multiline_data_are_not_split_for_detection(self):
        path = self.csv('"Preâmbulo\ncontinua"\n"Valor, líquido"\n"10,00\nobservado"\n')
        t = read_source(path, ReadOptions(header_row=2))[0]
        self.assertEqual(list(t.frame.columns), ['Valor, líquido'])
        self.assertEqual(t.frame.iloc[0, 0], '10,00\nobservado')
        self.assertEqual(t.row_numbers, [5])
        self.assertIsNone(t.source['delimiter'])

    def test_autodetection_never_hides_irregular_multicolumn_data(self):
        for sep in (',', ';', '|', '\t'):
            with self.subTest(sep=sep):
                path = self.csv(f'ID{sep}Valor\n1{sep}10\n2{sep}20{sep}SEGREDO\n')
                with self.assertRaises(InputError) as caught:
                    read_source(path)
                self.assertNotIn('SEGREDO', str(caught.exception))
                with self.assertRaises(InputError):
                    read_source(path, ReadOptions(sep=sep))

    def test_explicit_separator_stays_authoritative_for_single_column(self):
        path = self.csv('Valor\n10,00\n20,00\n')
        with self.assertRaises(InputError):
            read_source(path, ReadOptions(sep=','))
        self.assertEqual(read_source(path, ReadOptions(sep=';'))[0].frame['Valor'].tolist(),
                         ['10,00', '20,00'])

    def test_chartsheets_with_and_without_drawings_do_not_hide_data(self):
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        for drawing in (False, True):
            with self.subTest(drawing=drawing):
                wb = Workbook()
                ws = wb.active
                ws.title = 'Dados'
                ws.append(['ID', 'Valor'])
                ws.append(['001', '=1+1'])
                cs = wb.create_chartsheet('Grafico')
                if drawing:
                    chart = BarChart()
                    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=2), titles_from_data=True)
                    cs.add_chart(chart)
                path = self.root / 'synthetic.xlsx'
                wb.save(path)
                raw = path.read_bytes()
                tables = read_source(path)
                self.assertEqual([t.source['sheet'] for t in tables], ['Dados'])
                self.assertIsNone(tables[0].frame.iloc[0, 1])
                self.assertEqual(tables[0].frame.iloc[0, 0], '001')
                codes = {i['code'] for i in tables[0].issues}
                self.assertTrue({'non_tabular_sheets_skipped', 'formula_not_evaluated'} <= codes)
                self.assertEqual(path.read_bytes(), raw)
                with self.assertRaisesRegex(InputError, 'tabular'):
                    read_source(path, ReadOptions(sheet='Grafico'))

    def test_chart_only_workbook_is_refused_as_non_tabular(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.remove(wb.active)
        wb.create_chartsheet('Grafico')
        path = self.root / 'charts.xlsx'
        wb.save(path)
        with self.assertRaisesRegex(InputError, 'tabular'):
            read_source(path)


class RelationshipSafetyTests(unittest.TestCase):
    def test_nulls_do_not_invent_non_null_multiplication_in_either_direction(self):
        a, b = table({'ID': ['1', None]}), table({'ID': ['1', None, None]})
        a.id, b.id = 'a', 'b'
        for order in ([a, b], [b, a]):
            relations = candidates(order)['candidates']
            self.assertEqual(len(relations), 2)
            self.assertTrue(all(not r['multiplication_risk'] for r in relations))
            self.assertTrue(all(r['matched_join_rows'] == 1 for r in relations))
            self.assertEqual(sorted(r['null_source_rows'] for r in relations), [1, 2])

    def test_actual_multiplication_remains_directional(self):
        a, b = table({'ID': ['1', '1']}), table({'ID': ['1']})
        a.id, b.id = 'a', 'b'
        relations = {r['from_table']: r for r in candidates([a, b])['candidates']}
        self.assertFalse(relations['a']['multiplication_risk'])
        self.assertEqual(relations['a']['cardinality'], 'many_to_one')
        self.assertTrue(relations['b']['multiplication_risk'])
        self.assertEqual(relations['b']['cardinality'], 'one_to_many')


class AggregateRowTests(unittest.TestCase):
    def test_total_and_subtotals_are_hypotheses_with_bounded_row_evidence(self):
        for label in ('Total', 'Subtotal', 'TOTAL GERAL', 'Soma', 'Totais', 'Resultado'):
            with self.subTest(label=label):
                t = table({'ID': [label] * 25, 'Valor': ['100,00'] * 25})
                original = t.frame.copy(deep=True)
                p = profile_table(t, 'pt-BR')
                self.assertIn('footer_row_suspected', {i['code'] for i in p['issues']})
                event = next(i for i in p['issues'] if i['code'] == 'footer_row_suspected')
                self.assertEqual(event['severity'], 'WARNING')
                self.assertEqual(event['scope'], 'aggregation_review')
                self.assertEqual(event['evidence_kind'], 'inference')
                self.assertEqual(event['count'], 25)
                self.assertEqual(event['source_rows'], list(range(2, 22)))
                self.assertTrue(event['rows_truncated'])
                self.assertEqual(p['rows'], 25)
                self.assertEqual(p['grain']['status'], 'unknown')
                pd.testing.assert_frame_equal(t.frame, original)

    def test_ordinary_entities_and_labels_without_amount_are_not_totals(self):
        for data in ({'ID': ['001'], 'Nome': ['Total'], 'Valor': ['100']},
                     {'ID': ['Totalmente'], 'Valor': ['100']},
                     {'ID': ['Total'], 'Valor': [None]},
                     {'ID': ['Total'], 'Data': ['2026-01-01'], 'Valor': ['100']}):
            self.assertNotIn('footer_row_suspected', {i['code'] for i in profile_table(table(data))['issues']})

    def test_mixed_grain_prevents_cash_readiness_and_proposes_review_without_drop(self):
        t = table({'ID': ['001', 'Subtotal', '002', 'Total'],
                   'Valor': ['10', '10', '20', '30'], 'Moeda': ['BRL'] * 4,
                   'Tipo': ['Receita'] * 4, 'Status': ['Pago'] * 4,
                   'Data Pagamento': ['2026-01-01', None, '2026-01-02', None]})
        # Date columns on subtotal rows are absent; no detail row is discarded.
        p = profile_table(t)
        mapping = map_columns(t, p)
        r = financial.diagnose(t, p, mapping)
        for metric in r['indicators']:
            if metric['id'] != 'record_count':
                self.assertIn('detail_row_policy', metric['missing_requirements'])
                self.assertEqual(metric['status'], 'insufficient_evidence')
        steps = plan_transformations(t, p, mapping)
        review = next(s for s in steps if s['operation'] == 'review_aggregate_rows')
        self.assertFalse(review['executed'])
        self.assertTrue(review['preconditions'])
        self.assertEqual(len(t.frame), 4)
        self.assertNotIn('drop_rows', {s['operation'] for s in steps})

    def test_csv_and_xlsx_report_totals_without_altering_source_or_rows(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            csv_path = root / 'synthetic.csv'
            csv_path.write_text('ID;Valor\n1;10,00\nTotal;10,00\n', encoding='utf-8')
            wb = Workbook()
            wb.active.append(['ID', 'Valor'])
            wb.active.append(['1', '10,00'])
            wb.active.append(['Total', '10,00'])
            excel_path = root / 'synthetic.xlsx'
            wb.save(excel_path)
            for path in (csv_path, excel_path):
                with self.subTest(format=path.suffix):
                    raw = path.read_bytes()
                    report = analyze([path], locale='pt-BR')
                    self.assertEqual(report['tables'][0]['profile']['rows'], 2)
                    self.assertIn('footer_row_suspected', {i['code'] for i in report['issues']})
                    event = next(i for i in report['issues'] if i['code'] == 'footer_row_suspected')
                    self.assertEqual(event['source_rows'], [3])
                    self.assertEqual(path.read_bytes(), raw)
                    self.assertTrue(report['source_integrity'][0]['unchanged'])
