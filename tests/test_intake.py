import hashlib
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from powerbi_autopilot import intake, parsing


class ParsingTests(unittest.TestCase):
    def test_exact_regional_money_and_parentheses(self):
        for raw, locale, expected in [('R$ 1.234,56', None, '1234.56'),
                                      ('(1,234.56)', None, '-1234.56'),
                                      ('0,10', 'pt-BR', '0.10'),
                                      ('1.234', 'pt-BR', '1234'),
                                      ('(R$ 1.234,56)', None, '-1234.56'),
                                      ('-USD 20.00', None, '-20.00')]:
            with self.subTest(raw=raw):
                self.assertEqual(parsing.number(raw, locale).value, Decimal(expected))

    def test_ambiguous_or_malformed_money_is_not_guessed(self):
        for raw in ['1,234', '1.234']:
            self.assertEqual(parsing.number(raw).status, 'ambiguous')
        for raw in ['1.23.4', '12,34,56', 'NaN', 'Infinity', True, 'R$ banana', '-(20)']:
            self.assertEqual(parsing.number(raw).status, 'invalid', raw)
        self.assertEqual(parsing.number('').status, 'missing')
        self.assertEqual(parsing.number('1,234.56', 'pt-BR').status, 'invalid')

    def test_dates_refuse_ambiguous_and_serial_values(self):
        self.assertEqual(parsing.date('01/02/2026').status, 'ambiguous')
        self.assertEqual(parsing.date('01/02/2026', 'pt-BR').value.isoformat(), '2026-02-01')
        self.assertEqual(parsing.date('01/02/2026', 'en-US').value.isoformat(), '2026-01-02')
        self.assertEqual(parsing.date('2026-02-30').status, 'invalid')
        self.assertEqual(parsing.date(45000).status, 'invalid')
        self.assertEqual(parsing.date('13/02/2026').value.isoformat(), '2026-02-13')


class IntakeTests(unittest.TestCase):
    def test_header_formulas_are_blocked_without_evaluation(self):
        with tempfile.TemporaryDirectory() as directory:
            p = Path(directory)/'header.xlsx'
            wb = Workbook()
            wb.active.append(['=1+1', 'Valor'])
            wb.active.append(['001', '10.00'])
            wb.save(p)
            t = intake.read_source(p)[0]
            self.assertIn('formula_not_evaluated', {i['code'] for i in t.issues if i['severity'] == 'BLOCKER'})

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def test_csv_encoding_delimiter_headers_ids_and_raw_hash(self):
        p = self.root / 'raw.csv'
        raw = 'Código;Valor;Valor;\n001;1,50;2,50;\n002;NA;;\n'.encode('cp1252')
        p.write_bytes(raw)
        tables = intake.read_source(p)
        t = tables[0]
        self.assertEqual(list(t.frame.columns), ['Código', 'Valor', 'Valor__2', 'column_4'])
        self.assertEqual(t.frame.iloc[0, 0], '001')
        self.assertEqual(t.frame.iloc[1, 1], 'NA')
        self.assertEqual(t.row_numbers, [2, 3])
        self.assertEqual(t.source['sha256'], hashlib.sha256(raw).hexdigest())
        self.assertEqual(p.read_bytes(), raw)
        self.assertTrue({'encoding_fallback', 'duplicate_header', 'blank_header'} <= {i['code'] for i in t.issues})

    def test_workbook_all_sheets_formula_not_executed(self):
        p = self.root / 'book.xlsm'
        wb = Workbook()
        wb.active.title = 'Movimentos'
        wb.active.append(['ID', 'Valor'])
        wb.active.append(['001', '=1+2'])
        dim = wb.create_sheet('Clientes')
        dim.append(['ID', 'Nome'])
        dim.append(['001', 'Sintético'])
        wb.save(p)
        tables = intake.read_source(p)
        self.assertEqual([t.source['sheet'] for t in tables], ['Movimentos', 'Clientes'])
        self.assertIsNone(tables[0].frame.iloc[0, 1])
        self.assertIn('formula_not_evaluated', [i['code'] for i in tables[0].issues])
        self.assertEqual(len(intake.read_source(p, intake.ReadOptions(sheet='Clientes'))), 1)

    def test_bad_shapes_and_unsupported_formats_fail_without_data_leak(self):
        p = self.root / 'bad.csv'
        p.write_text('ID;Valor\n1;2;SEGREDO\n', encoding='utf-8')
        with self.assertRaisesRegex(intake.InputError, 'linha 2') as caught:
            intake.read_source(p, intake.ReadOptions(sep=';'))
        self.assertNotIn('SEGREDO', str(caught.exception))
        q = self.root / 'source.pbix'
        q.write_bytes(b'dummy')
        with self.assertRaisesRegex(intake.InputError, 'CSV, XLSX ou XLSM'):
            intake.read_source(q)

    def test_header_option_empty_csv_and_empty_sheet(self):
        p = self.root / 'preamble.csv'
        p.write_text('Relatório\nID;Valor\n1;10\n', encoding='utf-8')
        t = intake.read_source(p, intake.ReadOptions(header_row=2, sep=';'))[0]
        self.assertEqual(t.row_numbers, [3])
        p.write_bytes(b'')
        with self.assertRaises(intake.InputError):
            intake.read_source(p)
        p = self.root / 'empty.xlsx'
        Workbook().save(p)
        self.assertEqual(len(intake.read_source(p)[0].frame), 0)


if __name__ == '__main__':
    unittest.main()
