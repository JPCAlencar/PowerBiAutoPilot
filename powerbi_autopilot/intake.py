"""Local read-only intake. Workbook formulas and macros are never executed."""
import csv
from dataclasses import dataclass
import hashlib
import io
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import pandas as pd
from openpyxl.reader.excel import ExcelReader
from openpyxl.utils.exceptions import InvalidFileException

from .contracts import Table, issue


class InputError(ValueError):
    pass


@dataclass(frozen=True)
class ReadOptions:
    sheet: str | None = None
    sep: str | None = None
    encoding: str | None = None
    header_row: int = 1


class _TabularExcelReader(ExcelReader):
    """openpyxl 3.1 adapter: do not load drawings or non-tabular chartsheets.

    Filtering wb.sheetnames after load is too late: an empty chartsheet can
    already fail in openpyxl.read_chartsheet. This override also avoids reading
    chart/image payloads that are irrelevant to the tabular contract.
    """

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.skipped_chartsheets = []

    def read_chartsheet(self, sheet, rel):
        self.skipped_chartsheets.append(sheet.name)


def _csv_separator(text, header_row):
    """Only a separator observed outside quotes in the logical header qualifies.

    A multi-field header is never downgraded to one column to hide ragged rows.
    With no structural header separator, preserve a single field as a hypothesis.
    """
    structured, consistent = [], []
    for sep in ',;\t|':
        reader = csv.reader(io.StringIO(text, newline=''), delimiter=sep, strict=True)
        try:
            for _ in range(header_row):
                headers = next(reader)
            if len(headers) <= 1:
                continue
            structured.append(sep)
            if all(not record or len(record) == len(headers) for record in reader):
                consistent.append(sep)
        except (csv.Error, StopIteration):
            continue
    if len(consistent) == 1:
        return consistent[0]
    if structured:
        raise InputError('Delimitador incerto ou linhas irregulares; informe --sep.')
    return None


def _table(source, headers, records, row_numbers, diagnostics):
    table_id = source['sha256'][:16] + ':' + (source['sheet'] or 'csv')
    names, used, header_map = [], set(), []
    issues = [issue(code, severity, table_id, message, count=count) for code, severity, message, count in diagnostics]
    for i, raw in enumerate(headers, 1):
        base = str(raw).strip() if raw is not None else ''
        if not base:
            base = f'column_{i}'
            issues.append(issue('blank_header', 'WARNING', table_id, 'Cabeçalho vazio; nome posicional atribuído.', count=1))
        name, suffix = base, 2
        while name in used:
            name, suffix = f'{base}__{suffix}', suffix + 1
        if name != base:
            issues.append(issue('duplicate_header', 'WARNING', table_id, 'Cabeçalho repetido; nome posicional desambiguado.', column=name, count=1))
        names.append(name)
        used.add(name)
        header_map.append({'position': i, 'original': None if raw is None else str(raw), 'name': name})
    source['headers'] = header_map
    frame = pd.DataFrame(records, columns=names, dtype=object)
    return Table(table_id, frame, source, row_numbers, issues)


def read_source(path, options=None):
    options = options or ReadOptions()
    p = Path(path)
    if options.header_row < 1:
        raise InputError('header_row deve ser >= 1')
    if options.sep is not None and len(options.sep) != 1:
        raise InputError('Delimitador deve ter um caractere')
    if p.suffix.lower() not in {'.csv', '.xlsx', '.xlsm'}:
        raise InputError('Formato não suportado: use CSV, XLSX ou XLSM')
    try:
        raw = p.read_bytes()
    except OSError as exc:
        raise InputError('Não foi possível ler a fonte local.') from exc
    source = {'file': p.name, 'sha256': hashlib.sha256(raw).hexdigest(),
              'format': p.suffix.lower()[1:], 'sheet': None, 'header_row': options.header_row}
    if p.suffix.lower() == '.csv':
        if options.sheet is not None:
            raise InputError('--sheet só é aplicável a Excel')
        diagnostics = []
        encoding = options.encoding or 'utf-8-sig'
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError as exc:
            if options.encoding:
                raise InputError('Codificação informada não decodifica o arquivo.') from exc
            try:
                text, encoding = raw.decode('cp1252'), 'cp1252'
            except UnicodeDecodeError as fallback:
                raise InputError('Codificação desconhecida; informe --encoding.') from fallback
            diagnostics.append(('encoding_fallback', 'WARNING', 'UTF-8 inválido; cp1252 é hipótese. Confirme a codificação.', 1))
        except LookupError as exc:
            raise InputError('Codificação desconhecida.') from exc
        if '\x00' in text:
            raise InputError('CSV contém NUL; confirme a codificação (por exemplo utf-16).')
        if not text.strip():
            raise InputError('CSV vazio ou sem cabeçalho.')
        sep = options.sep
        if sep is None:
            sep = _csv_separator(text, options.header_row)
            if sep is None:
                diagnostics.append(('single_column_csv', 'INFO',
                                    'Cabeçalho sem separador estrutural; coluna única é hipótese. Confirme com --sep.', 1))
        source.update(encoding=encoding, delimiter=sep)
        # NUL is forbidden in the source, so it is a safe internal sentinel for
        # no delimiter while retaining standard CSV quoting/multiline semantics.
        reader = csv.reader(io.StringIO(text, newline=''), delimiter=sep or '\x00', strict=True)
        records, rows = [], []
        try:
            for _ in range(options.header_row):
                headers = next(reader)
            if not headers:
                raise InputError('Cabeçalho vazio.')
            for record in reader:
                row = reader.line_num
                if not record:
                    record = [''] * len(headers)
                if len(record) != len(headers):
                    raise InputError(f'Quantidade de campos incompatível na linha {row}.')
                records.append([None if value == '' else value for value in record])
                rows.append(row)
        except (csv.Error, StopIteration) as exc:
            raise InputError(f'CSV inválido próximo da linha {reader.line_num}.') from exc
        table = _table(source, headers, records, rows, diagnostics)
        for event in table.issues:
            if event['code'] == 'single_column_csv':
                event['evidence_kind'] = 'inference'
        return [table]
    reader = None
    try:
        reader = _TabularExcelReader(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            reader.read()
            wb = reader.wb
            if options.sheet in reader.skipped_chartsheets:
                raise InputError('Aba solicitada não é tabular.')
            if options.sheet is not None and options.sheet not in wb.sheetnames:
                raise InputError('Aba solicitada não encontrada.')
            tables = []
            for name in ([options.sheet] if options.sheet is not None else wb.sheetnames):
                ws = wb[name]
                cells = iter(ws.iter_rows(min_row=options.header_row))
                first = next(cells, ())
                headers = [None if c.data_type == 'f' else c.value for c in first]
                records, rows = [], []
                formula_rows = [options.header_row for c in first if c.data_type == 'f']
                for row_num, row in enumerate(cells, options.header_row+1):
                    values = []
                    for c in row:
                        if c.data_type == 'f':
                            formula_rows.append(row_num)
                            values.append(None)
                        else:
                            values.append(c.value)
                    records.append(values)
                    rows.append(row_num)
                table = _table({**source, 'sheet': name}, headers, records, rows, [])
                if formula_rows:
                    table.issues.append(issue('formula_not_evaluated', 'BLOCKER', table.id,
                                              'Fórmulas não avaliadas; células indisponíveis para análise.', rows=sorted(set(formula_rows)), count=len(formula_rows)))
                tables.append(table)
            if not tables:
                raise InputError('Nenhuma aba tabular disponível para análise.')
            if reader.skipped_chartsheets:
                tables[0].issues.append(issue('non_tabular_sheets_skipped', 'INFO', tables[0].id,
                                              'Abas de gráficos não tabulares ignoradas; dados tabulares preservados.',
                                              count=len(reader.skipped_chartsheets)))
            return tables
        finally:
            reader.archive.close()
    except InputError:
        raise
    except (OSError, ValueError, KeyError, BadZipFile, InvalidFileException, ParseError) as exc:
        raise InputError('Planilha inválida, protegida ou ilegível; use XLSX/XLSM válido.') from exc
    except Exception as exc:
        raise InputError('Falha interna na leitura tabular; diagnóstico não concluído.') from exc
